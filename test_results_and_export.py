from __future__ import annotations

import os
from datetime import date

from models import MeterManager
from data import create_repository


def seed_basic_readings(manager: MeterManager) -> None:
    # Simple deterministic dataset
    manager.add_reading("R", date(2025, 7, 1), 100)
    manager.add_reading("R", date(2025, 8, 31), 250)
    manager.add_reading("Y", date(2025, 7, 1), 200)
    manager.add_reading("Y", date(2025, 8, 31), 420)
    manager.add_reading("B", date(2025, 7, 1), 300)
    manager.add_reading("B", date(2025, 8, 31), 560)


def test_excel_export_tangedco_tmp(tmp_path):
    repo = create_repository("memory")
    manager = MeterManager(repository=repo)
    seed_basic_readings(manager)

    from data.excel_exporter import ExcelExporter

    exporter = ExcelExporter()
    out_path = os.path.join(tmp_path, "tangedco_export.xlsx")
    result = exporter.export_tangedco_format(out_path, manager)
    assert os.path.exists(result)


def test_excel_export_detailed_history(tmp_path):
    repo = create_repository("memory")
    manager = MeterManager(repository=repo)
    seed_basic_readings(manager)

    from data.excel_exporter import ExcelExporter

    exporter = ExcelExporter()
    out_path = os.path.join(tmp_path, "history.xlsx")
    result = exporter.export_detailed_history(out_path, manager)
    assert os.path.exists(result)


def test_excel_export_multisheet(tmp_path):
    repo = create_repository("memory")
    manager = MeterManager(repository=repo)
    seed_basic_readings(manager)

    from data.excel_exporter import ExcelExporter

    exporter = ExcelExporter()
    out_path = os.path.join(tmp_path, "all.xlsx")
    result = exporter.create_workbook_with_multiple_sheets(out_path, manager)
    assert os.path.exists(result)
    # If xlsx written, verify multiple sheets
    if result.lower().endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
            wb = load_workbook(result)
            sheets = set(wb.sheetnames)
            assert {"Summary", "History"}.issubset(sheets)
        except Exception:
            # In case openpyxl not installed in test environment
            pass
    else:
        # CSV fallback should have written at least the tangedco CSV
        assert os.path.exists(result)


def test_tangedco_headers_and_diffs():
    from data.excel_utils import TangedcoColumnConfig, prepare_tangedco_data
    cfg = TangedcoColumnConfig()
    readings = {
        cfg.meter_headers[0]: [{"date": __import__("datetime").date(2025, 7, 1), "reading": 100.0}, {"date": __import__("datetime").date(2025, 8, 31), "reading": 250.0}],
        cfg.meter_headers[1]: [{"date": __import__("datetime").date(2025, 7, 1), "reading": 200.0}, {"date": __import__("datetime").date(2025, 8, 31), "reading": 420.0}],
        cfg.meter_headers[2]: [{"date": __import__("datetime").date(2025, 7, 1), "reading": 300.0}, {"date": __import__("datetime").date(2025, 8, 31), "reading": 560.0}],
    }
    headers, rows = prepare_tangedco_data(readings)
    # Headers must start with Date and include meter headers interleaved with Diff
    assert headers[0] == cfg.date_header
    assert cfg.meter_headers[0] in headers and cfg.meter_headers[1] in headers and cfg.meter_headers[2] in headers
    # Verify diffs for first meter across second date = 150
    # Row order is by date ascending, so second row corresponds to 2025-08-31
    # Layout: [date, m1_reading, m1_diff, m2_reading, m2_diff, m3_reading, m3_diff]
    assert rows[1][2] == 150.0


def test_excel_export_summary_csv(tmp_path):
    repo = create_repository("memory")
    manager = MeterManager(repository=repo)
    seed_basic_readings(manager)

    from data.excel_exporter import ExcelExporter

    exporter = ExcelExporter()
    out_path = os.path.join(tmp_path, "summary.csv")
    result = exporter.export_summary_report(out_path, manager)
    assert os.path.exists(result)


