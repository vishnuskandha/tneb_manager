from __future__ import annotations

"""
Excel export functionality for readings, summaries, and detailed reports.

Exports are implemented primarily via pandas with openpyxl/xlsxwriter engines.
If pandas or Excel engines are unavailable, the implementation falls back to CSV
to ensure users can still export data.
"""

import os
from datetime import datetime, date
from typing import Any, Dict, List, Optional, Tuple


class ExcelExporter:
    """Handles creating Excel or CSV outputs based on available libraries."""

    def __init__(self) -> None:
        self._pandas = None
        self._openpyxl_available = False
        self._xlsxwriter_available = False
        try:
            import pandas as pd  # type: ignore

            self._pandas = pd
        except Exception:
            self._pandas = None
        try:
            import openpyxl  # noqa: F401  # type: ignore

            self._openpyxl_available = True
        except Exception:
            self._openpyxl_available = False
        try:
            import xlsxwriter  # noqa: F401  # type: ignore

            self._xlsxwriter_available = True
        except Exception:
            self._xlsxwriter_available = False

    # -------- Public API --------
    def export_tangedco_format(self, file_path: str, meter_manager: Any, start_date: Optional[date] = None, end_date: Optional[date] = None) -> str:
        """Export history in TANGEDCO layout. Returns output file path."""
        from data.excel_utils import (
            prepare_tangedco_data,
            validate_export_data,
            format_excel_date,
        )

        ok, err = validate_export_data(meter_manager)
        if not ok:
            raise RuntimeError(err or "Export validation failed")

        readings_by_meter = self._group_readings_by_tangedco_headers(meter_manager)
        # Optional billing-cycle filter
        if start_date and end_date:
            from data.excel_utils import calculate_billing_cycle_data
            for k, items in list(readings_by_meter.items()):
                readings_by_meter[k] = calculate_billing_cycle_data(items, start_date, end_date)
        headers, rows = prepare_tangedco_data(readings_by_meter)

        # Convert dates for Excel compliance
        for r in rows:
            if r and r[0] is not None:
                r[0] = format_excel_date(r[0])

        return self._write_tabular(file_path, headers, rows, sheet_name=self._default_sheet_name(start_date, end_date))

    def export_summary_report(self, file_path: str, meter_manager: Any) -> str:
        """Export overview summary including bills and recommendations."""
        summary = meter_manager.get_consumption_summary()
        rec = meter_manager.recommend_rotation()

        # Build a simple table for per-meter summary
        pm_units = summary.get("per_meter_units", {})
        pm_cost = summary.get("per_meter_cost", {})
        headers = ["Meter", "Units", "Cost"]
        rows: List[List[Any]] = []
        for meter_id in pm_units:
            rows.append([meter_id, float(pm_units[meter_id] or 0.0), float(pm_cost.get(meter_id, 0.0))])

        # Append total
        rows.append(["Total", float(summary.get("total_units", 0.0)), float(summary.get("total_cost", 0.0))])

        path = self._write_tabular(file_path, headers, rows, sheet_name="Summary")

        # If Excel is available, also add a second sheet with rotation details
        if self._pandas is not None and path.lower().endswith(".xlsx"):
            try:
                import pandas as pd  # type: ignore

                with pd.ExcelWriter(path, mode="a", engine=("openpyxl" if self._openpyxl_available else None)) as writer:  # type: ignore[arg-type]
                    rec_headers = ["Approaching", "Suggestions", "CombinedBill", "SeparateBill", "PotentialSavings"]
                    rec_rows = [[
                        ", ".join(rec.get("approaching", [])),
                        "; ".join(rec.get("suggestions", [])),
                        float(rec.get("economics", {}).get("combined_bill", 0.0)),
                        float(rec.get("economics", {}).get("separate_bill", 0.0)),
                        float(rec.get("economics", {}).get("potential_savings", 0.0)),
                    ]]
                    df_rec = pd.DataFrame(rec_rows, columns=rec_headers)
                    df_rec.to_excel(writer, sheet_name="Rotation", index=False)
            except Exception:
                # If appending fails, ignore; primary summary already exported
                pass

        return path

    def export_detailed_history(self, file_path: str, meter_manager: Any, start_date: Optional[date] = None, end_date: Optional[date] = None) -> str:
        """Export complete reading history for all meters."""
        history = meter_manager.get_readings_history()
        headers = ["Meter", "Date", "Reading"]
        rows: List[List[Any]] = []
        for r in history:
            if hasattr(r, 'meter_id'):
                m = r.meter_id
                d = r.reading_date
                val = r.current_reading
            else:
                m = r.get('meter_id') or r.get('meter')
                d = r.get('reading_date') or r.get('date')
                if isinstance(d, str):
                    from datetime import datetime as _dt
                    try:
                        d = _dt.fromisoformat(d).date()
                    except Exception:
                        pass
                val = r.get('current_reading') or r.get('reading')
            if start_date and end_date:
                try:
                    d_comp = d
                    if isinstance(d_comp, datetime):
                        d_comp = d_comp.date()
                    if d_comp is None or not (start_date <= d_comp <= end_date):
                        continue
                except Exception:
                    pass
            from data.excel_utils import format_excel_date
            if d is not None:
                d = format_excel_date(d)
            rows.append([m, d, val])

        return self._write_tabular(file_path, headers, rows, sheet_name="History")

    def create_workbook_with_multiple_sheets(self, file_path: str, meter_manager: Any, start_date: Optional[date] = None, end_date: Optional[date] = None) -> str:
        """Create multi-sheet workbook including TANGEDCO, Summary, and History."""
        # If pandas unavailable or not xlsx -> fall back to TANGEDCO-only export
        if self._pandas is None or not file_path.lower().endswith(".xlsx"):
            # Best-effort: write TANGEDCO CSV/Excel
            return self.export_tangedco_format(file_path, meter_manager, start_date, end_date)

        import pandas as pd  # type: ignore
        from data.excel_utils import prepare_tangedco_data, format_excel_date

        readings_by_meter = self._group_readings_by_tangedco_headers(meter_manager)
        if start_date and end_date:
            from data.excel_utils import calculate_billing_cycle_data
            for k, items in list(readings_by_meter.items()):
                readings_by_meter[k] = calculate_billing_cycle_data(items, start_date, end_date)

        # TANGEDCO sheet
        t_headers, t_rows = prepare_tangedco_data(readings_by_meter)
        for r in t_rows:
            if r[0] is not None:
                r[0] = format_excel_date(r[0])

        # Summary sheet
        summary = meter_manager.get_consumption_summary()
        pm_units = summary.get("per_meter_units", {})
        pm_cost = summary.get("per_meter_cost", {})
        s_headers = ["Meter", "Units", "Cost"]
        s_rows = [[m, float(pm_units[m] or 0.0), float(pm_cost.get(m, 0.0))] for m in pm_units]
        s_rows.append(["Total", float(summary.get("total_units", 0.0)), float(summary.get("total_cost", 0.0))])

        # History sheet
        h_headers = ["Meter", "Date", "Reading"]
        h_rows: List[List[Any]] = []
        # flatten readings_by_meter to rows
        for meter_name, entries in readings_by_meter.items():
            for r in entries:
                d = r.get("date")
                if d is not None and hasattr(d, "year"):
                    d = format_excel_date(d)
                h_rows.append([meter_name, d, r.get("reading")])

        engine = "openpyxl" if self._openpyxl_available else ("xlsxwriter" if self._xlsxwriter_available else None)
        if engine is None:
            # Fallback: write three CSV files with suffixes
            base, ext = os.path.splitext(file_path)
            self._write_csv(f"{base}_tangedco.csv", t_headers, t_rows)
            self._write_csv(f"{base}_summary.csv", s_headers, s_rows)
            self._write_csv(f"{base}_history.csv", h_headers, h_rows)
            return f"{base}_tangedco.csv"

        with pd.ExcelWriter(file_path, engine=engine) as writer:
            sheet_name = self._default_sheet_name(start_date, end_date)
            pd.DataFrame(t_rows, columns=t_headers).to_excel(writer, sheet_name=sheet_name, index=False)
            pd.DataFrame(s_rows, columns=s_headers).to_excel(writer, sheet_name="Summary", index=False)
            pd.DataFrame(h_rows, columns=h_headers).to_excel(writer, sheet_name="History", index=False)

        # Apply styling best-effort using openpyxl if available
        try:
            if self._openpyxl_available and file_path.lower().endswith(".xlsx"):
                from openpyxl import load_workbook
                from data.excel_utils import ExcelFormatHelper
                wb = load_workbook(file_path)
                try:
                    ws = wb[sheet_name]
                    ExcelFormatHelper.format_headers(ws, header_row_index=1)
                    last_row = ws.max_row
                    ExcelFormatHelper.apply_date_format(ws, f"A2:A{last_row}")
                except Exception:
                    pass
                try:
                    ws_sum = wb["Summary"]
                    last_row = ws_sum.max_row
                    ExcelFormatHelper.format_headers(ws_sum, header_row_index=1)
                    ExcelFormatHelper.apply_currency_format(ws_sum, f"C2:C{last_row}")
                except Exception:
                    pass
                wb.save(file_path)
        except Exception:
            pass

        return file_path

    # -------- Internal helpers --------
    def _default_sheet_name(self, start_date: Optional[date] = None, end_date: Optional[date] = None) -> str:
        import calendar
        now = datetime.now()
        if start_date and end_date:
            return f"{calendar.month_name[start_date.month]}-{calendar.month_name[end_date.month]} {end_date.strftime('%y')}"
        # Example: "July-August 25" style
        prev_month = 12 if now.month == 1 else now.month - 1
        prev_year = now.year - 1 if now.month == 1 else now.year
        _ = prev_year  # kept for potential future use
        return f"{calendar.month_name[prev_month]}-{calendar.month_name[now.month]} {now.strftime('%y')}"

    def _group_readings_by_tangedco_headers(self, meter_manager: Any) -> Dict[str, List[Dict[str, Any]]]:
        from data.excel_utils import TangedcoColumnConfig
        config = TangedcoColumnConfig()
        map_id_to_header = config.id_to_header
        res: Dict[str, List[Dict[str, Any]]] = {h: [] for h in config.meter_headers}
        for r in meter_manager.get_readings_history():
            if hasattr(r, 'meter_id'):
                m_id = getattr(r, 'meter_id', None)
                d = getattr(r, 'reading_date', None)
                val = getattr(r, 'current_reading', None)
            else:
                m_id = r.get('meter_id')
                d = r.get('reading_date') or r.get('date')
                if isinstance(d, str):
                    from datetime import datetime as _dt
                    try:
                        d = _dt.fromisoformat(d).date()
                    except Exception:
                        pass
                val = r.get('current_reading') or r.get('reading')
            key = map_id_to_header.get(m_id)
            if key is not None and d is not None and val is not None:
                try:
                    res[key].append({'date': d, 'reading': float(val)})
                except Exception:
                    pass
        for k in res:
            try:
                res[k].sort(key=lambda x: x['date'])
            except Exception:
                pass
        return res

    def _write_tabular(self, file_path: str, headers: List[str], rows: List[List[Any]], sheet_name: Optional[str] = None) -> str:
        if file_path.lower().endswith(".csv") or self._pandas is None:
            self._write_csv(file_path, headers, rows)
            return file_path
        engine = "openpyxl" if self._openpyxl_available else ("xlsxwriter" if self._xlsxwriter_available else None)
        if engine is None:
            # Fall back to CSV if no Excel engines are present
            base, _ = os.path.splitext(file_path)
            csv_path = f"{base}.csv"
            self._write_csv(csv_path, headers, rows)
            return csv_path
        import pandas as pd  # type: ignore

        with pd.ExcelWriter(file_path, engine=engine) as writer:
            df = pd.DataFrame(rows, columns=headers)
            df.to_excel(writer, index=False, sheet_name=(sheet_name or "Sheet1"))
        # Apply styling post-write (best-effort)
        try:
            if self._openpyxl_available and file_path.lower().endswith(".xlsx"):
                from openpyxl import load_workbook
                from data.excel_utils import ExcelFormatHelper
                wb = load_workbook(file_path)
                sh = (sheet_name or "Sheet1")
                try:
                    ws = wb[sh]
                    ExcelFormatHelper.format_headers(ws, header_row_index=1)
                    last_row = ws.max_row
                    # If Summary sheet, apply currency format to Cost column (C)
                    if sh == "Summary":
                        ExcelFormatHelper.apply_currency_format(ws, f"C2:C{last_row}")
                    else:
                        # Assume first column is date for other sheets
                        ExcelFormatHelper.apply_date_format(ws, f"A2:A{last_row}")
                except Exception:
                    pass
                wb.save(file_path)
        except Exception:
            pass
        return file_path

    def _write_csv(self, file_path: str, headers: List[str], rows: List[List[Any]]) -> None:
        import csv

        with open(file_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(headers)
            for r in rows:
                writer.writerow(r)


